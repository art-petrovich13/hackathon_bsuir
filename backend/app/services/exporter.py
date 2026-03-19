# backend/app/services/exporter.py
"""
Генерация DOCX отчёта по результатам сравнения НПА.
Использует python-docx (уже в requirements.txt как python-docx==1.1.2).
"""
from io import BytesIO
from datetime import datetime

from docx import Document as DocxDocument
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.models.comparison import Comparison
from app.models.diff_result import DiffResult


# Цвета для уровней риска
RISK_COLORS = {
    "CRITICAL": RGBColor(0xC0, 0x00, 0x00),
    "HIGH":     RGBColor(0xFF, 0x45, 0x00),
    "MEDIUM":   RGBColor(0xFF, 0xA5, 0x00),
    "LOW":      RGBColor(0x00, 0x80, 0x00),
}

RISK_LABELS = {
    "CRITICAL": "КРИТИЧЕСКИЙ",
    "HIGH":     "ВЫСОКИЙ",
    "MEDIUM":   "СРЕДНИЙ",
    "LOW":      "НИЗКИЙ",
}

SEMANTIC_LABELS = {
    "OBLIGATION_CHANGE": "Изменение обязательств",
    "SCOPE_CHANGE":      "Изменение сферы применения",
    "DEADLINE_CHANGE":   "Изменение сроков",
    "SUBJECT_CHANGE":    "Изменение субъекта",
    "SANCTION_CHANGE":   "Изменение ответственности",
    "COSMETIC":          "Косметические правки",
}

CHANGE_LABELS = {
    "ADDED":    "Добавлено",
    "DELETED":  "Удалено",
    "MODIFIED": "Изменено",
    "MOVED":    "Перемещено",
}


def _truncate(text: str | None, max_len: int = 100) -> str:
    if not text:
        return "—"
    return text if len(text) <= max_len else text[:max_len] + "..."


def generate_docx_report(
    comparison: Comparison,
    diff_results: list[DiffResult],
) -> bytes:
    """Генерирует DOCX отчёт и возвращает байты."""
    doc = DocxDocument()

    # Настройка шрифта по умолчанию
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)

    # ─── Титульный лист ───────────────────────────────────────────────────────
    title = doc.add_heading("Отчёт AI-анализатора НПА", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph(f"Дата: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub2 = doc.add_paragraph(f"ID сравнения: {comparison.id[:8]}...")
    sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("")

    # ─── Исполнительное резюме ────────────────────────────────────────────────
    doc.add_heading("Исполнительное резюме", level=1)

    summary = comparison.summary_json or {}
    total     = summary.get("total", len(diff_results))
    critical  = summary.get("critical", 0)
    high      = summary.get("high", 0)
    medium    = summary.get("medium", 0)
    low_count = summary.get("low", 0)
    risk_score = comparison.total_risk_score or 0.0

    resume = (
        f"В ходе анализа выявлено {total} изменений в нормативном акте. "
        f"Из них: критических — {critical}, высокого риска — {high}, "
        f"среднего риска — {medium}, низкого риска — {low_count}. "
        f"Общий индекс риска: {risk_score:.1f}/100."
    )
    if critical + high > 0:
        resume += f" Обнаружено {critical + high} изменений, требующих немедленного внимания юриста."

    doc.add_paragraph(resume)
    doc.add_paragraph("")

    # ─── Сводная таблица ─────────────────────────────────────────────────────
    doc.add_heading("Сводная таблица изменений", level=1)

    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"

    hdr = table.rows[0].cells
    for i, text in enumerate(["Пункт", "Тип", "Было", "Стало", "Риск"]):
        p = hdr[i].paragraphs[0]
        run = p.add_run(text)
        run.bold = True

    for result in diff_results:
        row = table.add_row()
        row.cells[0].text = result.section_path or ""
        row.cells[1].text = CHANGE_LABELS.get(result.change_type, result.change_type)
        row.cells[2].text = _truncate(result.old_text, 80)
        row.cells[3].text = _truncate(result.new_text, 80)

        risk = result.risk_level or "LOW"
        p = row.cells[4].paragraphs[0]
        run = p.add_run(RISK_LABELS.get(risk, risk))
        run.bold = True
        run.font.color.rgb = RISK_COLORS.get(risk, RGBColor(0, 0, 0))

    doc.add_paragraph("")

    # ─── Детальный анализ HIGH и CRITICAL ────────────────────────────────────
    high_critical = [r for r in diff_results if r.risk_level in ("HIGH", "CRITICAL")]
    if high_critical:
        doc.add_page_break()
        doc.add_heading("Детальный анализ критических изменений", level=1)

        for result in high_critical:
            risk = result.risk_level or "HIGH"
            h = doc.add_heading(
                f"Пункт {result.section_path}  [{RISK_LABELS.get(risk, risk)} РИСК]",
                level=2
            )
            for run in h.runs:
                run.font.color.rgb = RISK_COLORS.get(risk, RGBColor(0xFF, 0, 0))

            if result.old_text:
                p = doc.add_paragraph()
                p.add_run("Было: ").bold = True
                p.add_run(result.old_text)

            if result.new_text:
                p = doc.add_paragraph()
                p.add_run("Стало: ").bold = True
                p.add_run(result.new_text)

            if result.semantic_type:
                p = doc.add_paragraph()
                p.add_run("Тип изменения: ").bold = True
                p.add_run(SEMANTIC_LABELS.get(result.semantic_type, result.semantic_type))

            if result.law_reference:
                p = doc.add_paragraph()
                p.add_run("Нормативная основа: ").bold = True
                p.add_run(result.law_reference)

            if result.recommendation:
                p = doc.add_paragraph()
                p.add_run("Рекомендация: ").bold = True
                p.add_run(result.recommendation)

            if result.ai_confidence is not None:
                p = doc.add_paragraph()
                p.add_run("Уверенность AI: ").bold = True
                p.add_run(f"{result.ai_confidence * 100:.0f}%")

            doc.add_paragraph("")

    # ─── Дисклеймер ──────────────────────────────────────────────────────────
    doc.add_page_break()
    disc = doc.add_paragraph(
        "⚠ Данный отчёт носит информационный характер и не является юридической консультацией. "
        "Для принятия юридически значимых решений обратитесь к квалифицированному юристу."
    )
    disc.runs[0].font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    disc.runs[0].font.size = Pt(9)

    footer = doc.add_paragraph("Создано AI-ассистентом НПА-Анализатор · pravo.by")
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()